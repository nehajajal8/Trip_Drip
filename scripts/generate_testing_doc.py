"""
scripts/generate_testing_doc.py
Generates the Trip Drip Interactive Manual QA Testing & Execution Workbook (.docx).
Designed specifically for human testers to manually execute test cases, record actual outputs,
check PASS/FAIL status, and paste evidence screenshots into designated frames.
"""

import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# --- Color Palette Constants ---
HEX_PRIMARY_DARK = "1E3A8A"    # Deep Navy
HEX_PRIMARY_BLUE = "2563EB"    # Royal Blue
HEX_SECONDARY_TEAL = "0D9488"  # Dark Teal
HEX_TEXT_DARK = "1F2937"       # Charcoal
HEX_TEXT_MUTED = "4B5563"      # Slate Gray
HEX_BG_LIGHT = "F8FAFC"        # Light Slate Tint
HEX_BG_CALLOUT = "EFF6FF"      # Light Blue Tint
HEX_BORDER = "CBD5E1"          # Border Slate
HEX_BORDER_DASH = "94A3B8"     # Medium Slate for Screenshot Frames
HEX_ACCENT_AMBER = "D97706"    # Amber
HEX_SUCCESS_GREEN = "059669"   # Emerald Green

COLOR_PRIMARY_DARK = RGBColor(30, 58, 138)
COLOR_PRIMARY_BLUE = RGBColor(37, 99, 235)
COLOR_SECONDARY_TEAL = RGBColor(13, 148, 136)
COLOR_TEXT_DARK = RGBColor(31, 41, 55)
COLOR_TEXT_MUTED = RGBColor(75, 85, 99)
COLOR_SUCCESS_GREEN = RGBColor(5, 150, 105)

def set_cell_background(cell, hex_color):
    """Sets background color for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Sets inner margins (padding) in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)

def set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left=None, right=None, sz="4", val="single"):
    """Sets borders on a cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    borders_xml += f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{top}"/>' if top else '<w:top w:val="none"/>'
    borders_xml += f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{bottom}"/>' if bottom else '<w:bottom w:val="none"/>'
    borders_xml += f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{left}"/>' if left else '<w:left w:val="none"/>'
    borders_xml += f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{right}"/>' if right else '<w:right w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    tc_pr.append(parse_xml(borders_xml))

def format_table(table, col_widths=None, header_bg=HEX_PRIMARY_DARK):
    """Styles a standard data table with headers, borders, and alternating rows."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        is_header = (i == 0)
        bg_color = header_bg if is_header else (HEX_BG_LIGHT if i % 2 == 1 else "FFFFFF")
        for j, cell in enumerate(row.cells):
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=130, bottom=130, left=150, right=150)
            set_cell_borders(cell, top=HEX_BORDER, bottom=HEX_BORDER, sz="6" if is_header else "4")
            if col_widths and j < len(col_widths):
                cell.width = col_widths[j]
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    if is_header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(10)
                    else:
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = COLOR_TEXT_DARK

def add_heading_1(doc, text):
    """Adds a primary level-1 section heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY_DARK
    p_pr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="{HEX_PRIMARY_BLUE}"/></w:pBdr>')
    p_pr.append(pBdr)
    return p

def add_heading_2(doc, text):
    """Adds a level-2 sub-heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY_BLUE
    return p

def add_heading_3(doc, text):
    """Adds a level-3 sub-heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY_TEAL
    return p

def add_paragraph(doc, text, bold_prefix="", italic_prefix=""):
    """Adds a styled body paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_TEXT_DARK
    if italic_prefix:
        r_it = p.add_run(italic_prefix)
        r_it.font.name = "Calibri"
        r_it.font.size = Pt(10)
        r_it.font.italic = True
        r_it.font.color.rgb = COLOR_TEXT_MUTED
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = COLOR_TEXT_DARK
    return p

def add_bullet(doc, bold_title, text):
    """Adds a styled bullet item."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_bold = p.add_run(bold_title + ": ")
    r_bold.font.name = "Calibri"
    r_bold.font.size = Pt(10)
    r_bold.font.bold = True
    r_bold.font.color.rgb = COLOR_TEXT_DARK
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = COLOR_TEXT_DARK
    return p

def add_callout(doc, title, text, callout_type="NOTE"):
    """Adds a shaded callout box with colored left border."""
    border_color = HEX_PRIMARY_BLUE
    bg_color = HEX_BG_CALLOUT
    if callout_type == "IMPORTANT":
        border_color = HEX_ACCENT_AMBER
        bg_color = "FFFBEB"
    elif callout_type == "SUCCESS":
        border_color = HEX_SUCCESS_GREEN
        bg_color = "ECFDF5"

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.8)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    set_cell_borders(cell, top=None, bottom=None, left=border_color, right=None, sz="24")

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    r_icon = p.add_run(f"[{callout_type}] {title}: ")
    r_icon.font.name = "Calibri"
    r_icon.font.size = Pt(10)
    r_icon.font.bold = True
    r_icon.font.color.rgb = COLOR_SUCCESS_GREEN if callout_type == "SUCCESS" else (COLOR_PRIMARY_DARK if callout_type != "IMPORTANT" else RGBColor(180, 83, 9))

    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_TEXT_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_manual_test_card(doc, test_id, feature_name, page_url, pre_conds, steps, expected_output, screenshot_guidance):
    """Adds a complete interactive manual test case card with instructions, actual output lines, and screenshot placeholder."""
    add_heading_2(doc, f"{test_id}: {feature_name}")

    # Details table
    table = doc.add_table(rows=4, cols=2)
    row_headers = [
        ("Feature & URL", f"{feature_name} (Target: {page_url})"),
        ("Preconditions", pre_conds),
        ("Execution Steps", steps),
        ("Expected Output", expected_output)
    ]
    for idx, (label, val) in enumerate(row_headers):
        row = table.rows[idx]
        row.cells[0].paragraphs[0].add_run(label).font.bold = True
        row.cells[1].paragraphs[0].add_run(val)
    format_table(table, col_widths=[Inches(1.8), Inches(5.0)], header_bg=HEX_PRIMARY_BLUE)

    # Actual Results Block (Fillable by Tester)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    act_table = doc.add_table(rows=3, cols=2)
    act_headers = ["Tester Execution Field", "Tester Input / Recorded Observation"]
    for j, h in enumerate(act_headers):
        act_table.rows[0].cells[j].paragraphs[0].add_run(h)

    act_rows = [
        ("Test Verification Status", "[  ] PASS       [  ] FAIL       [  ] BLOCKED       Tested By: ______________   Date: ______________"),
        ("Actual Output Observed\n(Tester Notes & Remarks)", "Actual Result:\n____________________________________________________________________________________________________\n____________________________________________________________________________________________________")
    ]
    for idx, (label, val) in enumerate(act_rows):
        row = act_table.rows[idx + 1]
        row.cells[0].paragraphs[0].add_run(label).font.bold = True
        row.cells[1].paragraphs[0].add_run(val)
    format_table(act_table, col_widths=[Inches(2.2), Inches(4.6)], header_bg=HEX_SECONDARY_TEAL)

    # Screenshot Frame
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    s_table = doc.add_table(rows=1, cols=1)
    s_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_cell = s_table.rows[0].cells[0]
    s_cell.width = Inches(6.8)
    set_cell_background(s_cell, "F8FAFC")
    set_cell_margins(s_cell, top=240, bottom=240, left=200, right=200)
    set_cell_borders(s_cell, top=HEX_BORDER_DASH, bottom=HEX_BORDER_DASH, left=HEX_BORDER_DASH, right=HEX_BORDER_DASH, sz="12", val="dashed")

    s_p = s_cell.paragraphs[0]
    s_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_p.paragraph_format.space_before = Pt(3)
    s_p.paragraph_format.space_after = Pt(3)

    r_cam = s_p.add_run("📷 [PASTE TEST EVIDENCE SCREENSHOT HERE]\n")
    r_cam.font.name = "Calibri"
    r_cam.font.size = Pt(10.5)
    r_cam.font.bold = True
    r_cam.font.color.rgb = COLOR_PRIMARY_BLUE

    r_desc = s_p.add_run(f"Evidence for {test_id}: {screenshot_guidance}\n")
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(9.5)
    r_desc.font.bold = True
    r_desc.font.color.rgb = COLOR_TEXT_DARK

    r_tip = s_p.add_run("(Take screenshot with Win+Shift+S, click inside this box, and press Ctrl+V)")
    r_tip.font.name = "Calibri"
    r_tip.font.size = Pt(8.5)
    r_tip.font.italic = True
    r_tip.font.color.rgb = COLOR_TEXT_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def build_testing_document(output_path):
    print(f"Generating Interactive Manual QA Testing Workbook: {output_path}...")
    doc = Document()

    # --- Page Setup ---
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)

    # --- Header & Footer ---
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run("Trip Drip (India Edition) — Manual QA Testing Workbook & Execution Log")
    f_run.font.name = "Calibri"
    f_run.font.size = Pt(8.5)
    f_run.font.color.rgb = COLOR_TEXT_MUTED

    # --- Title Banner ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("TRIP DRIP — MANUAL QA TESTING WORKBOOK")
    t_run.font.name = "Calibri"
    t_run.font.size = Pt(22)
    t_run.font.bold = True
    t_run.font.color.rgb = COLOR_PRIMARY_DARK

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    s_run = sub_p.add_run("Hands-On Feature Test Cases, Expected Outputs, Actual Output Recording Fields & Screenshot Evidence Frames")
    s_run.font.name = "Calibri"
    s_run.font.size = Pt(12)
    s_run.font.italic = True
    s_run.font.color.rgb = COLOR_PRIMARY_BLUE

    # --- Metadata Table ---
    meta_table = doc.add_table(rows=6, cols=2)
    meta_data = [
        ("Workbook Identifier", "QA-TD-2026-MANUAL-WB"),
        ("Project Under Test", "Trip Drip (India Edition v2.4)"),
        ("Document Purpose", "Interactive QA Workbook for Manual Testing, Recording Observations & Pasting Screenshots"),
        ("Lead Tester / User", "Neha Jajal / Assigned QA Engineer"),
        ("Application Test URL", "http://localhost:5173 (Vite Dev Server) & Backend API: http://localhost:3001"),
        ("Execution Status", "In Progress — Complete test cases below, mark Pass/Fail, and insert screenshots")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        row.cells[0].paragraphs[0].add_run(label).font.bold = True
        row.cells[1].paragraphs[0].add_run(val)
    format_table(meta_table, col_widths=[Inches(2.3), Inches(4.5)], header_bg=HEX_PRIMARY_DARK)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ─────────────────────────────────────────────────────────────────
    # SECTION 1: MANUAL TESTING INSTRUCTIONS & ENVIRONMENT SETUP
    # ─────────────────────────────────────────────────────────────────
    add_heading_1(doc, "1. Testing Instructions & How to Use This Workbook")
    add_paragraph(doc, "This workbook is designed for you to manually test the core features of Trip Drip in your browser, verify that each feature works as expected, record your actual observations, and paste evidence screenshots.")

    add_heading_2(doc, "1.1 How to Start the App for Testing")
    add_bullet(doc, "Step 1", "Open terminal in project directory: `c:\\Users\\Neha Jajal\\Downloads\\trip-drip (2)\\trip-drip\\trip-drip`.")
    add_bullet(doc, "Step 2", "Start the frontend dev server: run `npm run dev`. Verify the app is running at `http://localhost:5173`.")
    add_bullet(doc, "Step 3", "In a second terminal, start the API backend (optional): run `npm run dev:api`.")
    add_bullet(doc, "Step 4", "Open `http://localhost:5173` in Google Chrome, Microsoft Edge, or Firefox.")

    add_heading_2(doc, "1.2 How to Record Test Results & Insert Screenshots")
    add_bullet(doc, "Recording Status", "In each test card below, check [X] PASS if the behavior matches Expected Output, or [X] FAIL if an error occurs.")
    add_bullet(doc, "Writing Observations", "Type your actual notes, such as load time, UI responsiveness, or specific values displayed.")
    add_bullet(doc, "Inserting Screenshots", "Press `Win + Shift + S` on your Windows keyboard, capture the active page/modal, click inside the designated dashed box in this Word document, and press `Ctrl + V`.")

    add_callout(doc, "Recommended Tools for Testing", 
                "1. Browser: Google Chrome or Microsoft Edge (DevTools available via F12).\n"
                "2. Screenshot Utility: Windows Snipping Tool (`Win + Shift + S`) or Lightshot.\n"
                "3. Mobile Viewport Testing: Press F12 in browser, click Device Emulation icon (Ctrl+Shift+M) to test responsive UI.",
                "TIP")

    # ─────────────────────────────────────────────────────────────────
    # SECTION 2: MANUAL FEATURE TEST CASES & SCREENSHOT WORKBOOK
    # ─────────────────────────────────────────────────────────────────
    add_heading_1(doc, "2. Feature-by-Feature Test Cases & Evidence Worksheets")

    # TEST CASE 1: Trip Creation Wizard
    add_manual_test_card(
        doc,
        test_id="TC-01",
        feature_name="5-Step Trip Planning Wizard & Itinerary Generation",
        page_url="http://localhost:5173/trips/new",
        pre_conds="App is running. User is logged in or creates a guest trip. Navigate to '/trips/new'.",
        steps="1. Click '+ New Trip' button on dashboard or navigate to '/trips/new'.\n"
              "2. Step 1: Select Destination (e.g. 'Jaipur' or 'Mumbai').\n"
              "3. Step 2: Pick Start and End dates for a 5-day duration (e.g. Oct 1 to Oct 5).\n"
              "4. Step 3: Set Total Budget: ₹25,000 and Group Size: 4.\n"
              "5. Step 4: Choose Transport Preference (Train) and Style (Heritage & Culture).\n"
              "6. Step 5: Click 'Generate Trip' button and observe generation speed.",
        expected_output="1. Progress bar advances smoothly across all 5 wizard steps.\n"
                        "2. Destination image card reflects chosen city.\n"
                        "3. Budget splits preview displays estimated stay, travel, and food allowances.\n"
                        "4. System creates trip and redirects to Trip Dashboard (`/trips/:id`) in < 2 seconds.\n"
                        "5. Generated itinerary displays exactly 5 distinct days with unique routes.",
        screenshot_guidance="Capture the New Trip Wizard (Step 3 or Step 4) showing destination, dates, and budget inputs."
    )

    # TEST CASE 2: Day-by-Day Itinerary Schedule
    add_manual_test_card(
        doc,
        test_id="TC-02",
        feature_name="Trip Dashboard & Deduplicated Day Schedule",
        page_url="http://localhost:5173/trips/:id",
        pre_conds="Active trip opened in Trip Dashboard.",
        steps="1. Open any generated trip (e.g. Jaipur or Mumbai).\n"
              "2. Examine the Day 1, Day 2, Day 3, Day 4, and Day 5 tabs in the schedule.\n"
              "3. Check the route text on Day 1 vs Day 2 vs Day 3.\n"
              "4. Click on 'Day 2' tab; inspect activities, theme name, and wardrobe advice.\n"
              "5. Verify weather widget card displays temperature and packing tips.",
        expected_output="1. Header card displays destination title, dates, countdown, and budget balance.\n"
                        "2. Day 1 route is COMPLETELY DIFFERENT from Day 2 and Day 3 (Zero duplicate days).\n"
                        "3. Each day has a unique cultural theme (e.g. Forts & Palaces, Bazaars, Culinary).\n"
                        "4. Wardrobe card suggests specific climate-appropriate attire for that day.\n"
                        "5. Interactive Leaflet map renders destination location marker.",
        screenshot_guidance="Capture the Trip Dashboard showing Day 1 and Day 2 schedule tabs with unique routes."
    )

    # TEST CASE 3: Hisaab-Kitaab Expense Splitting & UPI QR
    add_manual_test_card(
        doc,
        test_id="TC-03",
        feature_name="Hisaab-Kitaab Expense Resolution & UPI Scan-to-Pay",
        page_url="http://localhost:5173/trips/:id/expenses",
        pre_conds="Active trip opened. Squad members defined (e.g. Aarav, Priya, Rohan, Sneha).",
        steps="1. Click 'Expenses' in the trip navigation bar.\n"
              "2. Click '+ Add Expense'; enter Label: 'SUV Highway Rental', Amount: ₹4000, Paid By: Aarav.\n"
              "3. Add second expense: Label: 'Dhaba Lunch', Amount: ₹2000, Paid By: Priya.\n"
              "4. Scroll down to 'Hisaab-Kitaab (Settlements)' section.\n"
              "5. Verify calculated settlements (who owes whom) and click 'Pay via UPI QR' button.",
        expected_output="1. Total Group Spend accurately sums to ₹6,000.\n"
                        "2. Per-person fair share accurately displays ₹1,500.\n"
                        "3. Minimal Cashflow algorithm generates exactly 3 clean settlement transactions.\n"
                        "4. Clicking 'Pay via UPI QR' opens modal rendering dynamic UPI QR code.\n"
                        "5. Mobile link shows valid 'upi://pay?pa=...&cu=INR' scheme.",
        screenshot_guidance="Capture the Hisaab-Kitaab Settlement summary cards showing who owes whom and the UPI QR Code modal."
    )

    # TEST CASE 4: Chai & Station Radar
    add_manual_test_card(
        doc,
        test_id="TC-04",
        feature_name="Chai & Station Radar (Transit Gastronomy Modal)",
        page_url="http://localhost:5173/trips/:id (Click 'Chai & Station Radar')",
        pre_conds="Trip Dashboard is loaded. Click the 'Chai & Station Radar' quick-action button.",
        steps="1. Locate the 'Chai & Station Radar' button on the dashboard and click it.\n"
              "2. In the radar modal, view the transit corridor (e.g. CSMT Mumbai or Lonavala Western Ghats).\n"
              "3. Verify iconic tapris listed (e.g. Cutting Masala Chai, Aram Vada Pav, Chikki).\n"
              "4. Switch to a different corridor tab (e.g. Old Delhi or Jaipur Junction).\n"
              "5. Check the 'Rail Neer Verified Water Checkpoint' badge.",
        expected_output="1. Modal opens smoothly with glassmorphism dark/light backdrop.\n"
                        "2. Lists verified tapri stalls, platform locations, and opening hours.\n"
                        "3. Displays average prices in INR (e.g. ₹15 Cutting Chai, ₹25 Vada Pav).\n"
                        "4. Clean water verification badge shows 'Rail Neer Verified Platform 1'.\n"
                        "5. Modal closes cleanly on 'Close' or outside backdrop click.",
        screenshot_guidance="Capture the Chai & Station Radar modal showing CSMT Mumbai or Lonavala tapris and food items."
    )

    # TEST CASE 5: Smart Wardrobe & Suitability Scoring
    add_manual_test_card(
        doc,
        test_id="TC-05",
        feature_name="Smart Wardrobe Upload & Weather Suitability",
        page_url="http://localhost:5173/trips/:id/wardrobe",
        pre_conds="Active trip opened. Navigate to 'Wardrobe' tab.",
        steps="1. Click 'Wardrobe' in the top navigation.\n"
              "2. Click '+ Add Garment' or upload an apparel photo (shirt, jacket, shoes).\n"
              "3. Select category: 'Tops / Shirts' and color: 'White'.\n"
              "4. Observe client-side image compression in console / UI.\n"
              "5. Click Save and inspect the garment card for suitability rating (0 - 10).",
        expected_output="1. Image optimizes instantly (< 200ms) without UI freezing.\n"
                        "2. Uploaded item appears in the trip wardrobe grid.\n"
                        "3. Card displays category badge, color tag, and destination suitability score.\n"
                        "4. Dynamic badge indicates weather compatibility (e.g. '9.2/10 - Ideal for warm humidity').",
        screenshot_guidance="Capture the Wardrobe grid displaying uploaded apparel items and suitability rating badges."
    )

    # TEST CASE 6: Cross-Brand Fashion Comparison & Cart
    add_manual_test_card(
        doc,
        test_id="TC-06",
        feature_name="Cross-Brand Fashion Comparison & Dynamic Cart",
        page_url="http://localhost:5173/trips/:id/shop",
        pre_conds="Active trip opened. Navigate to 'Shop' tab.",
        steps="1. Click 'Shop' tab in top navigation.\n"
              "2. View the 4-brand side-by-side comparison table (Zara vs H&M vs Uniqlo vs Westside).\n"
              "3. Filter by category (e.g. 'Linen Tops' or 'Cargo Trousers').\n"
              "4. Click 'Add to Cart' on an in-budget item (e.g. Uniqlo AIRism or H&M Relaxed Linen).\n"
              "5. Observe the shopping cart badge count and budget deduction indicator.",
        expected_output="1. Brand benchmark displays accurate comparative pricing and fabric specifications.\n"
                        "2. 'In Budget' green badge appears on items within trip financial limit.\n"
                        "3. Adding item increments shopping cart count in real time.\n"
                        "4. Estimated trip remaining budget recalculates with cart items deducted.",
        screenshot_guidance="Capture the Cross-Brand Comparison page showing Zara, H&M, Uniqlo, and Westside product cards."
    )

    # TEST CASE 7: Travel Journal & Memory Log
    add_manual_test_card(
        doc,
        test_id="TC-07",
        feature_name="Travel Journal & Photo Memories",
        page_url="http://localhost:5173/trips/:id/journal",
        pre_conds="Active trip opened. Navigate to 'Journal' tab.",
        steps="1. Click 'Journal' in top navigation.\n"
              "2. Click '+ Add Entry'.\n"
              "3. Select today's date and enter note: 'Watched sunset from Hawa Mahal, amazing kulhad chai!'.\n"
              "4. Attach a sample travel photo.\n"
              "5. Click 'Save Memory' and verify entry displays in the timeline.",
        expected_output="1. New journal card renders with timestamp, photo thumbnail, and reflection text.\n"
                        "2. Entry persists locally in IndexedDB `journal` store.\n"
                        "3. Refreshing the browser page retains the entry without data loss.",
        screenshot_guidance="Capture the Journal timeline showing the newly logged travel memory and attached photo."
    )

    # TEST CASE 8: AI Travel Butler (Trip Concierge Action Bus)
    add_manual_test_card(
        doc,
        test_id="TC-08",
        feature_name="AI Travel Butler Floating Concierge & Tool Actions",
        page_url="http://localhost:5173/trips/:id (Floating Chat Icon)",
        pre_conds="Any trip sub-page opened. Look for floating blue chat icon at bottom-right.",
        steps="1. Click the floating chat bubble icon at bottom-right of screen.\n"
              "2. Verify drawer slides open with greeting personalized to destination.\n"
              "3. Type prompt: 'Can you add Zara linen shirt to my cart?' and press Enter.\n"
              "4. Type second prompt: 'Show me chai radar highway dhabas' and press Enter.\n"
              "5. Verify AI responds and dispatches client action badges.",
        expected_output="1. Chat drawer expands smoothly without obscuring core page navigation.\n"
                        "2. Concierge identifies 'ADD_TO_CART' intent and triggers `trip_cart_updated`.\n"
                        "3. Cart badge count increments immediately.\n"
                        "4. Second query automatically triggers and launches the Chai Radar Modal.\n"
                        "5. Full conversational history persists across sub-page transitions.",
        screenshot_guidance="Capture the expanded AI Concierge drawer showing user prompt and assistant action response."
    )

    # TEST CASE 9: Offline Mode PWA & Offline Banner
    add_manual_test_card(
        doc,
        test_id="TC-09",
        feature_name="Zero-Connectivity Offline Resilience & PWA Caching",
        page_url="http://localhost:5173/trips/:id (Simulated Offline)",
        pre_conds="Trip loaded with itinerary. Browser DevTools open (press F12).",
        steps="1. Press F12 in browser, go to 'Network' tab, change Throttling to 'Offline'.\n"
              "2. Observe the top of the app interface for the Offline status alert.\n"
              "3. Navigate to Trip Dashboard day schedule tabs.\n"
              "4. Navigate to Expenses tab and view existing settlements.\n"
              "5. Restore network to 'Online' in DevTools.",
        expected_output="1. Persistent orange banner appears: 'You are currently offline. Viewing cached travel data'.\n"
                        "2. Day-by-Day schedule and itinerary remain 100% visible and interactive.\n"
                        "3. Expenses and Hisaab-Kitaab remain accessible via IndexedDB cache.\n"
                        "4. Re-enabling network seamlessly clears the offline banner.",
        screenshot_guidance="Capture the browser showing the 'Offline' warning banner while viewing the cached trip itinerary."
    )

    # TEST CASE 10: Cascade Trip Deletion
    add_manual_test_card(
        doc,
        test_id="TC-10",
        feature_name="Trip Deletion & Cascade Storage Cleanup",
        page_url="http://localhost:5173/trips",
        pre_conds="At least one test trip exists in the trips list.",
        steps="1. Navigate to `/trips` gallery.\n"
              "2. Locate the test trip card and click the Delete / Trash icon.\n"
              "3. Verify confirmation modal appears warning of permanent deletion.\n"
              "4. Confirm deletion.\n"
              "5. Check dashboard: trip card is gone; check IndexedDB (Application tab in F12): cart, expenses, journal for that trip ID are wiped.",
        expected_output="1. Confirmation modal prevents accidental deletion.\n"
                        "2. Upon confirmation, trip card smoothly unmounts from gallery.\n"
                        "3. Cascade cleanup ensures zero orphaned cart, expense, or journal entries in local storage.\n"
                        "4. UI toast or alert confirms 'Trip deleted successfully'.",
        screenshot_guidance="Capture the Trip Deletion confirmation modal or the updated trips gallery after deletion."
    )

    # ─────────────────────────────────────────────────────────────────
    # SECTION 3: MASTER TEST EXECUTION SUMMARY & SIGN-OFF SHEET
    # ─────────────────────────────────────────────────────────────────
    add_heading_1(doc, "3. Master Test Execution Summary & Sign-Off Sheet")
    add_paragraph(doc, "After completing the 10 manual test cases above, tally your results in the summary table below and record the final release verdict:")

    summary_table = doc.add_table(rows=12, cols=4)
    sum_headers = ["Test ID", "Feature Tested", "Your Verified Status", "Defect ID (If Failed)"]
    for j, h in enumerate(sum_headers):
        summary_table.rows[0].cells[j].paragraphs[0].add_run(h)

    tests_list = [
        ("TC-01", "5-Step Trip Planning Wizard & Itinerary Gen", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TC-02", "Trip Dashboard & Deduplicated Schedule", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TC-03", "Hisaab-Kitaab Splitter & UPI QR Modal", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TC-04", "Chai & Station Radar Transit Modal", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TC-05", "Smart Wardrobe Upload & Suitability", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TC-06", "Cross-Brand Fashion Comparison & Cart", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TC-07", "Travel Journal & Photo Memories", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TC-08", "AI Travel Butler Concierge & Action Bus", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TC-09", "Zero-Connectivity Offline Resilience", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TC-10", "Trip Deletion & Cascade Storage Cleanup", "[  ] PASS   [  ] FAIL", "None / ________"),
        ("TOTALS", "10 Feature Suites Tested", "Pass: ___ / 10 | Fail: ___", "Overall: [  ] READY")
    ]
    for idx, row_data in enumerate(tests_list):
        row = summary_table.rows[idx + 1]
        for col_idx, val in enumerate(row_data):
            run = row.cells[col_idx].paragraphs[0].add_run(val)
            if idx == len(tests_list) - 1:
                run.font.bold = True
    format_table(summary_table, col_widths=[Inches(1.2), Inches(3.0), Inches(1.8), Inches(1.2)])

    # Sign-off card
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_heading_2(doc, "3.1 Tester Sign-Off & Verification Verdict")
    sign_table = doc.add_table(rows=4, cols=2)
    sign_rows = [
        ("Tester Name & Signature", "Name: _______________________________   Signature: _______________________________"),
        ("Date of Manual Testing", "Date: _______________________________   Location / Device: ________________________"),
        ("Final Quality Verdict", "[  ] APPROVED FOR RELEASE / PRODUCTION        [  ] REJECTED — BLOCKING DEFECTS FOUND"),
        ("General Comments / Feedback", "_____________________________________________________________________________________\n_____________________________________________________________________________________")
    ]
    for idx, (label, val) in enumerate(sign_rows):
        row = sign_table.rows[idx]
        row.cells[0].paragraphs[0].add_run(label).font.bold = True
        row.cells[1].paragraphs[0].add_run(val)
    format_table(sign_table, col_widths=[Inches(2.4), Inches(4.4)], header_bg=HEX_PRIMARY_DARK)

    doc.save(output_path)
    print(f"Successfully generated Manual QA Testing Workbook: {output_path}")

if __name__ == "__main__":
    os.makedirs("docs", exist_ok=True)
    root_out = "Trip_Drip_Testing_Document.docx"
    docs_out = os.path.join("docs", "Trip_Drip_Testing_Document.docx")
    build_testing_document(root_out)
    shutil.copy(root_out, docs_out)
    print(f"Copied to {docs_out}")
